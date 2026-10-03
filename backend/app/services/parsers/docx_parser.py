"""
DOCX parser for extracting text, tables, and images from Word documents.
"""

import io
import logging
from pathlib import Path
from typing import List, Optional

from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

from app.services.parsers.base import (
    BaseParser, ParsedDocument, ExtractedText, ExtractedImage, ExtractedTable
)

logger = logging.getLogger(__name__)


class DOCXParser(BaseParser):
    """Parser for Microsoft Word (.docx) files."""

    def supported_extensions(self) -> List[str]:
        return [".docx"]

    async def parse(self, file_path: Path, document_id: str, output_dir: Path) -> ParsedDocument:
        logger.info(f"Parsing DOCX: {file_path.name}")

        result = ParsedDocument(
            document_id=document_id,
            filename=file_path.name,
            file_type="docx",
        )

        images_dir = output_dir / "images"
        images_dir.mkdir(parents=True, exist_ok=True)

        try:
            doc = DocxDocument(str(file_path))

            # Extract paragraphs
            current_heading = None
            page_estimate = 1  # DOCX doesn't have real page numbers
            char_count = 0

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # Track headings
                if para.style.name.startswith("Heading"):
                    current_heading = text

                char_count += len(text)
                # Rough page estimate: ~3000 chars per page
                page_estimate = max(1, char_count // 3000 + 1)

                result.texts.append(ExtractedText(
                    text=text,
                    page_number=page_estimate,
                    heading=current_heading,
                    metadata={
                        "style": para.style.name,
                        "is_heading": para.style.name.startswith("Heading"),
                    }
                ))

            result.page_count = page_estimate

            # Extract tables
            for table_idx, table in enumerate(doc.tables):
                try:
                    rows_data = []
                    headers = []

                    for row_idx, row in enumerate(table.rows):
                        cells = [cell.text.strip() for cell in row.cells]
                        if row_idx == 0:
                            headers = cells
                        else:
                            rows_data.append(cells)

                    if headers or rows_data:
                        result.tables.append(ExtractedTable(
                            headers=headers,
                            rows=rows_data,
                            page_number=None,
                            metadata={
                                "table_index": table_idx,
                                "extraction_method": "python-docx"
                            }
                        ))
                except Exception as e:
                    logger.warning(f"Failed to extract table {table_idx}: {e}")

            # Extract images
            img_idx = 0
            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    try:
                        image_part = rel.target_part
                        image_bytes = image_part.blob
                        content_type = image_part.content_type
                        ext = content_type.split("/")[-1] if content_type else "png"
                        if ext == "jpeg":
                            ext = "jpg"

                        image_filename = f"{document_id}_img{img_idx}.{ext}"
                        image_path = images_dir / image_filename

                        with open(image_path, "wb") as f:
                            f.write(image_bytes)

                        # Get dimensions
                        try:
                            img = Image.open(io.BytesIO(image_bytes))
                            width, height = img.size
                        except Exception:
                            width, height = None, None

                        result.images.append(ExtractedImage(
                            image_path=str(image_path),
                            image_bytes=image_bytes,
                            image_format=ext,
                            width=width,
                            height=height,
                            metadata={"content_type": content_type}
                        ))
                        img_idx += 1
                    except Exception as e:
                        logger.warning(f"Failed to extract image: {e}")

            # Document metadata
            core_props = doc.core_properties
            result.metadata = {
                "title": core_props.title or "",
                "author": core_props.author or "",
                "subject": core_props.subject or "",
                "created": str(core_props.created) if core_props.created else "",
                "modified": str(core_props.modified) if core_props.modified else "",
                "file_size_bytes": file_path.stat().st_size,
            }

        except Exception as e:
            logger.error(f"Error parsing DOCX {file_path.name}: {e}")
            result.errors.append(f"DOCX parsing error: {str(e)}")

        logger.info(
            f"DOCX parsed: ~{result.page_count} pages, "
            f"{len(result.texts)} text blocks, "
            f"{len(result.images)} images, "
            f"{len(result.tables)} tables"
        )
        return result
