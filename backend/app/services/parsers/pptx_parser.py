"""
PPTX parser for extracting text, images, and tables from PowerPoint files.
"""

import io
import logging
from pathlib import Path
from typing import List

try:
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:
    Presentation = None
    Inches = None
from PIL import Image

from app.services.parsers.base import (
    BaseParser, ParsedDocument, ExtractedText, ExtractedImage, ExtractedTable
)

logger = logging.getLogger(__name__)


class PPTXParser(BaseParser):
    """Parser for Microsoft PowerPoint (.pptx) files."""

    def supported_extensions(self) -> List[str]:
        return [".pptx"]

    async def parse(self, file_path: Path, document_id: str, output_dir: Path) -> ParsedDocument:
        logger.info(f"Parsing PPTX: {file_path.name}")

        result = ParsedDocument(
            document_id=document_id,
            filename=file_path.name,
            file_type="pptx",
        )

        images_dir = output_dir / "images"
        images_dir.mkdir(parents=True, exist_ok=True)

        if Presentation is None:
            result.errors.append("python-pptx package not installed")
            return result

        try:
            prs = Presentation(str(file_path))
            result.page_count = len(prs.slides)

            for slide_num, slide in enumerate(prs.slides, 1):
                slide_texts = []
                slide_title = None

                for shape in slide.shapes:
                    # Extract text from text frames
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text.strip()
                            if text:
                                slide_texts.append(text)

                        # Detect slide title
                        if shape.shape_type is not None:
                            try:
                                if hasattr(shape, "placeholder_format") and shape.placeholder_format:
                                    if shape.placeholder_format.idx == 0:  # Title placeholder
                                        slide_title = shape.text_frame.text.strip()
                            except Exception:
                                pass

                    # Extract tables
                    if shape.has_table:
                        try:
                            table = shape.table
                            headers = []
                            rows_data = []

                            for row_idx, row in enumerate(table.rows):
                                cells = [cell.text.strip() for cell in row.cells]
                                if row_idx == 0:
                                    headers = cells
                                else:
                                    rows_data.append(cells)

                            if headers or rows_data:
                                result.tables.append(ExtractedTable(
                                    headers=headers,
                                    rows=rows_data,
                                    page_number=slide_num,
                                    metadata={
                                        "slide_number": slide_num,
                                        "extraction_method": "python-pptx"
                                    }
                                ))
                        except Exception as e:
                            logger.warning(f"Failed to extract table from slide {slide_num}: {e}")

                    # Extract images
                    if shape.shape_type == 13:  # Picture
                        try:
                            image = shape.image
                            image_bytes = image.blob
                            ext = image.content_type.split("/")[-1] if image.content_type else "png"
                            if ext == "jpeg":
                                ext = "jpg"

                            image_filename = f"{document_id}_slide{slide_num}_img.{ext}"
                            image_path = images_dir / image_filename

                            with open(image_path, "wb") as f:
                                f.write(image_bytes)

                            try:
                                img = Image.open(io.BytesIO(image_bytes))
                                width, height = img.size
                            except Exception:
                                width, height = None, None

                            result.images.append(ExtractedImage(
                                image_path=str(image_path),
                                image_bytes=image_bytes,
                                image_format=ext,
                                page_number=slide_num,
                                width=width,
                                height=height,
                                metadata={
                                    "slide_number": slide_num,
                                    "content_type": image.content_type,
                                }
                            ))
                        except Exception as e:
                            logger.warning(f"Failed to extract image from slide {slide_num}: {e}")

                # Combine slide text
                if slide_texts:
                    full_text = "\n".join(slide_texts)
                    result.texts.append(ExtractedText(
                        text=full_text,
                        page_number=slide_num,
                        heading=slide_title,
                        metadata={
                            "slide_number": slide_num,
                            "slide_title": slide_title,
                        }
                    ))

            # Metadata
            result.metadata = {
                "slide_count": len(prs.slides),
                "slide_width": str(prs.slide_width),
                "slide_height": str(prs.slide_height),
                "file_size_bytes": file_path.stat().st_size,
            }

        except Exception as e:
            logger.error(f"Error parsing PPTX {file_path.name}: {e}")
            result.errors.append(f"PPTX parsing error: {str(e)}")

        logger.info(
            f"PPTX parsed: {result.page_count} slides, "
            f"{len(result.texts)} text blocks, "
            f"{len(result.images)} images, "
            f"{len(result.tables)} tables"
        )
        return result
